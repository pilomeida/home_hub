"""One-off historical import: reads 'Utilities/Electricity/Electricity
consumption.xlsx' (the user's own reconciled analysis) directly into
Document + Transaction + UtilityReading triples, matching each row to its
source PDF by invoice number (extracted from the PDF text, since filenames
don't encode it). No LLM calls — the spreadsheet is already-verified ground
truth, preserving its manual corrections (a fees-box fix, one excluded
credit note) rather than re-deriving them from the PDFs.

Not part of the reviewed application code (see the utilities-electricity
spec's "Out of Scope" section).

Usage:
    python scripts/backfill_electricity_history.py
"""

import re
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from sqlmodel import Session

from app.db import engine
from app.models.document import Document, DocumentSource, DocumentStatus
from app.models.transaction import Category, Transaction, TransactionType
from app.models.utility_reading import UtilityReading, UtilityType
from app.services.dedup import find_existing_document_by_hash
from app.services.storage import save_upload

ELECTRICITY_DIR = Path(__file__).resolve().parent.parent / "Utilities" / "Electricity"
XLSX_PATH = ELECTRICITY_DIR / "Electricity consumption.xlsx"
PROVIDER = "Coopérnico"

_MONTH_ABBR = {
    "Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6,
    "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12,
}


def _parse_month_label(month_str: str) -> str:
    """'Nov 2024' -> '2024-11'"""
    abbr, year = month_str.split()
    return f"{year}-{_MONTH_ABBR[abbr]:02d}"


def _parse_billing_period(period_str: str) -> tuple[date, date]:
    """'26/10/2024 – 25/11/2024' -> (date(2024,10,26), date(2024,11,25))"""
    dates = re.findall(r"(\d{2})/(\d{2})/(\d{4})", period_str)
    if len(dates) != 2:
        raise ValueError(f"Could not parse billing period: {period_str!r}")
    (d1, m1, y1), (d2, m2, y2) = dates
    return date(int(y1), int(m1), int(d1)), date(int(y2), int(m2), int(d2))


def _extract_invoice_number(pdf_path: Path) -> str | None:
    """Returns the invoice number from a 'Fatura: X' line, or None if this
    PDF is a credit note ('Nota Crédito: X') or has no matching line."""
    text = subprocess.run(
        ["pdftotext", str(pdf_path), "-"], capture_output=True, text=True, check=True
    ).stdout
    match = re.search(r"^Fatura:\s*(.+)$", text, re.MULTILINE)
    if match:
        return match.group(1).strip()
    return None


def _build_invoice_map() -> dict[str, Path]:
    """invoice_number -> pdf path. Credit notes and unmatched PDFs are
    skipped (reported separately). The one known duplicate PDF (same
    invoice number, re-downloaded) resolves to whichever filename sorts
    first, deterministically."""
    invoice_map: dict[str, Path] = {}
    skipped: list[str] = []
    for pdf_path in sorted(ELECTRICITY_DIR.glob("*.pdf")):
        invoice_number = _extract_invoice_number(pdf_path)
        if invoice_number is None:
            skipped.append(pdf_path.name)
            continue
        if invoice_number in invoice_map:
            print(
                f"  (duplicate invoice {invoice_number}: keeping "
                f"{invoice_map[invoice_number].name}, ignoring {pdf_path.name})"
            )
            continue
        invoice_map[invoice_number] = pdf_path

    if skipped:
        print(f"Skipped {len(skipped)} PDF(s) with no 'Fatura:' line (credit notes): {skipped}")
    return invoice_map


def _read_rows() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Data"]
    rows = []
    for row in ws.iter_rows(min_row=5, max_row=25, values_only=True):
        month, period, invoice_number = row[0], row[1], row[2]
        if month is None or month == "TOTAL":
            continue
        if invoice_number is None:
            print(f"  Skipping {month}: no invoice on file")
            continue
        rows.append({
            "month": month,
            "period": period,
            "invoice_number": invoice_number,
            "consumption_kwh": row[3],
            "cost_total": row[4],
            "cost_per_kwh": row[5],
            "energy": row[6],
            "power": row[7],
            "fees_taxes": row[8],
            "vat": row[9],
        })
    return rows


def backfill() -> None:
    print(f"Reading {XLSX_PATH}")
    rows = _read_rows()
    print(f"Found {len(rows)} billed rows in the spreadsheet\n")

    print("Extracting invoice numbers from PDFs...")
    invoice_map = _build_invoice_map()
    print(f"Matched {len(invoice_map)} unique invoice(s) from PDFs\n")

    created = 0
    skipped_dup = 0
    errors = []

    with Session(engine) as session:
        for row in rows:
            invoice_number = row["invoice_number"]
            pdf_path = invoice_map.get(invoice_number)
            if pdf_path is None:
                errors.append(f"{row['month']}: no PDF found for invoice {invoice_number!r}")
                continue

            content = pdf_path.read_bytes()
            file_path, content_hash = save_upload(pdf_path.name, content)

            existing = find_existing_document_by_hash(session, content_hash)
            if existing is not None:
                print(f"SKIP (duplicate)  {row['month']} ({invoice_number}) -> document #{existing.id}")
                skipped_dup += 1
                continue

            period_label = _parse_month_label(row["month"])
            billing_start, billing_end = _parse_billing_period(row["period"])

            document = Document(
                filename=pdf_path.name,
                file_path=file_path,
                content_hash=content_hash,
                source=DocumentSource.MANUAL,
                status=DocumentStatus.PROCESSED,
                doc_type="bill",
                uploaded_by="backfill-electricity-history-script",
            )
            session.add(document)
            session.commit()
            session.refresh(document)

            transaction = Transaction(
                document_id=document.id,
                provider=PROVIDER,
                category=Category.ELECTRICITY,
                transaction_type=TransactionType.DEBIT,
                amount=row["cost_total"],
                currency="EUR",
                due_date=None,
                paid_date=billing_end,
                statement_period=period_label,
            )
            session.add(transaction)

            reading = UtilityReading(
                document_id=document.id,
                utility_type=UtilityType.ELECTRICITY,
                period_label=period_label,
                billing_period_start=billing_start,
                billing_period_end=billing_end,
                invoice_number=invoice_number,
                consumption_value=row["consumption_kwh"],
                consumption_unit="kWh",
                cost_total=row["cost_total"],
                cost_per_unit=row["cost_per_kwh"],
                energy_cost=row["energy"],
                power_cost=row["power"],
                fees_taxes_cost=row["fees_taxes"],
                vat_cost=row["vat"],
            )
            session.add(reading)
            session.commit()

            print(
                f"OK  {row['month']} ({invoice_number}) -> document #{document.id}, "
                f"{row['consumption_kwh']} kWh, {row['cost_total']:.2f} EUR"
            )
            created += 1

    print("\n--- Summary ---")
    print(f"Created:         {created}")
    print(f"Skipped (dup):   {skipped_dup}")
    print(f"Errors:          {len(errors)}")
    for e in errors:
        print(f"  - {e}")


if __name__ == "__main__":
    backfill()
